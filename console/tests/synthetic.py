"""
가짜 이름으로 만든 기록지 — 저장소가 공개라 실제 선수 기록지는 올리지 않는다.
실제 기록지와 같은 시트 구조와, 검증기가 잡아야 하는 실제 결함을 그대로 재현한다.

  · 시즌 시트: 타자 블록 → 투수 블록이 위아래로 쌓인 구조 + '통산' 합계 행
  · 학기 시트: "99 Spring", "99 Fall" (+ "2099 Spring 시즌" 합계 행)
  · 경기별 시트: "99 Spring (타자)" — "7. 가선수" 라벨 블록, 출전하지 않은 경기(전부 0), 빈 템플릿 "0. 선수"
  · 결함: 합계 행 SUM 범위 누락(마지막 선수), 합계 행 이닝을 표기끼리 더함(10.2+5.2=15.4),
          학기 합 ≠ 시즌 시트(1명), 같은 이름 다른 배번(1명), 통산 시트 이닝 표기 오류
"""
from __future__ import annotations

from pathlib import Path

import openpyxl

SEASON = "2099"
YY = "99"

BAT_HEADER = ["배번", "이름", "경기", "타석", "타수", "득점", "안타", "2루타", "3루타", "홈런", "타점", "볼넷", "사구", "삼진", "도루",
              "타율", "출루율", "장타율", "OPS"]
PIT_HEADER = ["배번", "이름", "경기", "승", "패", "세", "홀", "이닝", "피안타", "실점", "자책", "4구", "사구", "삼진", "피홈런", "타자", "타수",
              "ERA", "WHIP", "피안타율"]
GAME_BAT_HEADER = ["날짜", "상대팀", *BAT_HEADER[2:]]
GAME_PIT_HEADER = ["날짜", "상대팀", *PIT_HEADER[2:]]

# (배번, 이름, 학기) → 경기별 타격 [경기, 타석, 타수, 득점, 안타, 2루타, 3루타, 홈런, 타점, 볼넷, 사구, 삼진, 도루]
GAMES = [("2099-04-10", "가상팀A", "Spring"), ("2099-05-08", "가상팀B", "Spring"), ("2099-10-02", "가상팀C", "Fall")]
BATTING = {
    (7, "가선수"): [[1, 4, 3, 1, 2, 1, 0, 0, 1, 1, 0, 0, 1], [1, 3, 3, 0, 1, 0, 0, 0, 0, 0, 0, 1, 0], [1, 4, 4, 2, 2, 0, 0, 1, 3, 0, 0, 1, 1]],
    (11, "나선수"): [[1, 4, 2, 1, 1, 0, 0, 0, 0, 2, 0, 0, 2], [0] * 13, [1, 3, 3, 0, 0, 0, 0, 0, 0, 0, 0, 2, 0]],
    (23, "다선수"): [[1, 3, 3, 0, 1, 0, 1, 0, 2, 0, 0, 1, 0], [1, 4, 3, 1, 1, 0, 0, 0, 0, 1, 0, 1, 1], [1, 3, 2, 1, 1, 0, 0, 0, 1, 1, 0, 0, 0]],
    (35, "라선수"): [[0] * 13, [1, 2, 2, 0, 0, 0, 0, 0, 0, 0, 0, 2, 0], [1, 4, 3, 2, 2, 1, 0, 0, 1, 1, 0, 0, 3]],
}
# (배번, 이름) → 경기별 투구 [경기, 승, 패, 세, 홀, 이닝, 피안타, 실점, 자책, 4구, 사구, 삼진, 피홈런, 타자, 타수]
PITCHING = {
    (11, "나선수"): [[1, 1, 0, 0, 0, 5.1, 4, 2, 1, 2, 0, 6, 0, 22, 19], [0] * 15, [1, 0, 1, 0, 0, 5.1, 6, 4, 3, 3, 1, 4, 1, 26, 21]],
    (35, "라선수"): [[1, 0, 0, 0, 1, 2.2, 2, 1, 1, 1, 0, 3, 0, 11, 10], [1, 1, 0, 0, 0, 3.0, 1, 0, 0, 0, 0, 5, 0, 10, 10], [0] * 15],
}


def _add_ip(a: float, b: float) -> float:
    outs = round(a) * 3 + round((a % 1) * 10) + round(b) * 3 + round((b % 1) * 10)
    full, rem = divmod(outs, 3)
    return float(f"{full}.{rem}")


def _sum_rows(rows, ip_col=None):
    out = [0] * len(rows[0])
    for r in rows:
        for i, v in enumerate(r):
            out[i] = _add_ip(out[i], v) if i == ip_col else out[i] + v
    return out


def _bat_rates(r):
    g, pa, ab, runs, h, d, t, hr, rbi, bb, hbp, so, sb = r
    avg = h / ab if ab else 0
    obp = (h + bb + hbp) / pa if pa else 0
    slg = (h + d + 2 * t + 3 * hr) / ab if ab else 0
    return [avg, obp, slg, obp + slg]


def _season_block_rows(term: str | None, *, drop_last_from_sum=False, notation_sum_in_footer=False):
    idx = [i for i, (_, _, t) in enumerate(GAMES) if term is None or t == term]
    bat = []
    for (num, name), games in BATTING.items():
        tot = _sum_rows([games[i] for i in idx])
        if tot[0]:
            bat.append([num, name, *tot, *_bat_rates(tot)])
    pit = []
    for (num, name), games in PITCHING.items():
        tot = _sum_rows([games[i] for i in idx], ip_col=5)
        if tot[0]:
            pit.append([num, name, *tot, 0, 0, 0])
    footer_bat_rows = bat[:-1] if drop_last_from_sum else bat
    fb = _sum_rows([r[2:15] for r in footer_bat_rows])
    if notation_sum_in_footer:
        fp = [sum(r[2 + i] for r in pit) for i in range(15)]  # 기록지 SUM: 이닝 표기를 소수로 더함
    else:
        fp = _sum_rows([r[2:17] for r in pit], ip_col=5)
    return bat, pit, fb, fp


def _write_summary(ws, label, term, **kw):
    bat, pit, fb, fp = _season_block_rows(term, **kw)
    ws.append(["타자기록"])
    ws.append(BAT_HEADER)
    for r in bat:
        ws.append(r)
    ws.append([label, None, *fb, *_bat_rates(fb)])
    ws.append([])
    ws.append(["투수기록"])
    ws.append(PIT_HEADER)
    for r in pit:
        ws.append(r)
    ws.append([label, None, *fp, 0, 0, 0])


def build_season_file(folder: Path) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = SEASON
    _write_summary(ws, "통산", None, notation_sum_in_footer=True)
    _write_summary(wb.create_sheet(f"{YY} Spring"), f"{SEASON} Spring 시즌", "Spring", drop_last_from_sum=True)
    _write_summary(wb.create_sheet(f"{YY} Fall"), f"{SEASON} Fall 시즌", "Fall")

    # 학기 합 ≠ 시즌 시트: 시즌 시트의 다선수 도루를 1 늘린다
    for row in ws.iter_rows(min_row=3):
        if row[1].value == "다선수":
            row[14].value += 1
            break

    for term in ("Spring", "Fall"):
        idx = [i for i, (_, _, t) in enumerate(GAMES) if t == term]
        wsb = wb.create_sheet(f"{YY} {term} (타자)")
        for (num, name), games in BATTING.items():
            # 같은 이름 다른 배번: Fall 경기별 시트의 라선수를 #36으로 적음
            shown = 36 if (term == "Fall" and name == "라선수") else num
            wsb.append([f"{shown}. {name}"])
            wsb.append(GAME_BAT_HEADER)
            for i in idx:
                date, opp, _ = GAMES[i]
                wsb.append([date.replace("-", "."), opp, *games[i], *_bat_rates(games[i])])
            wsb.append([])
        wsb.append(["0. 선수"])
        wsb.append(GAME_BAT_HEADER)
        wsb.append([GAMES[idx[0]][0].replace("-", "."), "", *([0] * 13)])

        wsp = wb.create_sheet(f"{YY} {term} (투수)")
        for (num, name), games in PITCHING.items():
            wsp.append([f"{num}. {name}"])
            wsp.append(GAME_PIT_HEADER)
            for i in idx:
                date, opp, _ = GAMES[i]
                wsp.append([date.replace("-", "."), opp, *games[i], 0, 0, 0])
            wsp.append([])

    path = folder / f"{SEASON}_Utah_Devils_Player_Stat.xlsx"
    wb.save(path)
    return path


def build_career_file(folder: Path) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Career Totals"
    _write_summary(ws, "통산", None)
    # 통산 시트 이닝 표기 오류 (실제 기록지의 85.4 같은 값)
    for row in ws.iter_rows():
        if row[1].value == "나선수" and isinstance(row[7].value, float):
            row[7].value = 10.4
    _write_summary(wb.create_sheet(SEASON), "통산", None)
    wb.create_sheet("Hall of Fame").append(["명예의 전당"])
    path = folder / "Career_Total_Utah_Devils_Player_Stat.xlsx"
    wb.save(path)
    return path


def build_all(folder: Path) -> list[Path]:
    folder.mkdir(parents=True, exist_ok=True)
    return [build_season_file(folder), build_career_file(folder)]
