"""기록지 업로드 화면의 '경기 기록지 v3.0' 부분."""
from __future__ import annotations

import hashlib
import io

import openpyxl
import pandas as pd
import streamlit as st

from .diffing import diff
from .game_sheet import is_game_workbook, parse_game_workbook
from .session import actor, refresh
from .stats import add_batting_rates

GAME_COMPARE_B = ["pa", "ab", "hits", "bb", "so", "runs", "rbi", "sb"]
GAME_COMPARE_P = ["outs", "er", "ha", "bb", "so"]


class UnreadableWorkbook(ValueError):
    """엑셀로 열 수 없는 파일 (손상, 다른 형식, 암호 등)."""


def is_game_file(uploaded) -> bool:
    """경기 기록지(타자 기록·투수 기록 시트)인지. 엑셀로 열 수 없으면 UnreadableWorkbook."""
    try:
        wb = openpyxl.load_workbook(io.BytesIO(uploaded.getvalue()), read_only=True)
    except Exception as e:  # openpyxl은 zipfile.BadZipFile, KeyError, InvalidFileException 등 다양하게 던진다
        raise UnreadableWorkbook(f"{uploaded.name}: 엑셀 파일로 열 수 없습니다 ({type(e).__name__}). "
                                 "Excel에서 .xlsx 또는 .xlsm으로 다시 저장해 올려 주세요.") from e
    try:
        return is_game_workbook(wb.sheetnames)
    finally:
        wb.close()


@st.cache_data(show_spinner="경기 기록지 분석 중…")
def _parse(digest: str, name: str, raw: bytes):
    return parse_game_workbook(io.BytesIO(raw), name)


def _existing_rows(data, info):
    """같은 시즌·날짜·상대 경기가 이미 저장돼 있으면 그 경기의 기록 (선수 이름 포함)."""
    games = data["games"]
    if games is None or games.empty or not info.get("date"):
        return None, pd.DataFrame(), pd.DataFrame()
    g = games[(games["season"].astype(str) == str(info["season"])) & (games["date"].astype(str) == info["date"])
              & (games["opponent"].astype(str) == str(info["opponent"]))]
    if g.empty:
        return None, pd.DataFrame(), pd.DataFrame()
    gid = g.iloc[0]["id"]
    names = data["players"].set_index("id")["name"]

    def rows(raw):
        if raw is None or raw.empty or "game_id" not in raw:
            return pd.DataFrame()
        part = raw[raw["game_id"] == gid].copy()
        part["name"] = part["player_id"].map(names)
        part["number"] = 0
        return part

    bat, pit = rows(data["bat_raw"]), rows(data["pit_raw"])
    if not pit.empty:
        from .stats import ip_to_outs
        pit["outs"] = pit["ip"].map(lambda v: ip_to_outs(v, strict=False))
    return gid, bat, pit


def render_game_files(files, repo, data):
    st.header("경기 기록지")
    v3_ready = repo.read_only or repo.v3_ready()
    if not v3_ready:
        st.error("DB에 경기 기록지 v3.0 칸(희생플라이·희생번트·포수방해·투구수)과 경기 교체 함수가 아직 없습니다. "
                 "저장소의 `sql/005_game_sheet_v3.sql`을 Supabase SQL Editor에서 적용하면 저장할 수 있습니다. "
                 "아래 미리보기는 그대로 볼 수 있습니다.", icon=":material/database_off:")

    for f in files:
        try:
            sheet = _parse(hashlib.sha1(f.getvalue()).hexdigest(), f.name, f.getvalue())
        except Exception as e:
            st.error(f"{f.name}: 경기 기록지를 읽지 못했습니다 — {e}", icon=":material/error:")
            continue
        info = sheet.info
        st.subheader(f"{info.get('date') or '날짜 없음'} vs {info.get('opponent') or '상대 없음'}")
        st.caption(f"{f.name} · 기록지 {info.get('version') or '버전 미표기'}")

        c = st.columns(5)
        c[0].metric("점수", f"{info.get('score_us', '-')} : {info.get('score_them', '-')}")
        c[1].metric("홈/원정", {True: "홈", False: "원정"}.get(info.get("is_home"), "-"))
        c[2].metric("타자", len(sheet.batting))
        c[3].metric("투수", len(sheet.pitching))
        c[4].metric("검수", sheet.review_status or "미확인")

        if sheet.warnings:
            st.warning("확인이 필요한 항목\n\n" + "\n".join(f"- {w}" for w in sheet.warnings), icon=":material/fact_check:")
        else:
            st.success("검수 완료, 대조 항목 모두 일치", icon=":material/check_circle:")

        bat = add_batting_rates(sheet.batting.assign(g=1))
        st.dataframe(bat[["order", "name", "position", "pa", "ab", "hits", "doubles", "triples", "hr", "bb", "hbp", "so",
                          "sf", "sh", "ci", "runs", "rbi", "sb", "avg", "obp"]],
                     hide_index=True, width="stretch",
                     column_config={"order": "타순", "name": "이름", "position": "포지션", "pa": "타석", "ab": "타수",
                                    "hits": "안타", "doubles": "2루타", "triples": "3루타", "hr": "홈런", "bb": "볼넷",
                                    "hbp": "사구", "so": "삼진", "sf": "희생플라이", "sh": "희생번트", "ci": "포수방해",
                                    "runs": "득점", "rbi": "타점", "sb": "도루",
                                    "avg": st.column_config.NumberColumn("타율", format="%.3f"),
                                    "obp": st.column_config.NumberColumn("출루율", format="%.3f")})
        st.dataframe(sheet.pitching[["name", "decision", "ip", "outs", "pitches", "ha", "runs_allowed", "er", "bb", "hbp",
                                     "so", "hr_allowed"]],
                     hide_index=True, width="stretch",
                     column_config={"name": "이름", "decision": "결과", "ip": "이닝", "outs": "아웃", "pitches": "투구수",
                                    "ha": "피안타", "runs_allowed": "실점", "er": "자책", "bb": "볼넷", "hbp": "사구",
                                    "so": "삼진", "hr_allowed": "피홈런"})

        gid, old_b, old_p = _existing_rows(data, info)
        if gid is not None:
            changes = pd.concat([
                diff(sheet.batting.assign(number=0), old_b, GAME_COMPARE_B).assign(구분="타자"),
                diff(sheet.pitching.assign(number=0), old_p, GAME_COMPARE_P).assign(구분="투수"),
            ])
            if changes.empty:
                st.info(f"이미 저장된 경기 #{gid}와 기록이 같습니다.", icon=":material/check:")
            else:
                st.markdown(f"**이미 저장된 경기 #{gid}와 다른 점** — 저장하면 이 경기 기록이 교체됩니다.")
                st.dataframe(changes[["구분", "상태", "이름", "변경 내용"]], hide_index=True, width="stretch")

        if repo.read_only:
            st.info("데모 모드에서는 저장할 수 없습니다.", icon=":material/lock:")
            continue
        if not v3_ready:
            continue

        # 기록지에는 배번이 없어 이름으로 선수를 찾는다 — 동명이인이면 저장할 수 없다
        ambiguous = repo.ambiguous_names(list(sheet.batting["name"]) + list(sheet.pitching["name"]))
        if ambiguous:
            st.error("동명이인 때문에 이 경기는 저장할 수 없습니다.\n\n" + "\n\n".join(f"- {m}" for m in ambiguous),
                     icon=":material/person_off:")
            continue

        key = hashlib.sha1(f.getvalue()).hexdigest()[:8]
        ok = st.checkbox("위 내용을 확인했고 이 경기 기록을 저장합니다." + ("" if sheet.is_reviewed else " (검수 미완료)"),
                         key=f"game_ok_{key}")
        if st.button("경기 저장", type="primary", disabled=not ok, key=f"game_save_{key}"):
            res = repo.replace_game(sheet, actor())
            if res.ok:
                refresh()
                st.success(f"{res.message} — " + ", ".join(f"{k} {v}" for k, v in (res.counts or {}).items()))
            else:
                st.error(res.message)
