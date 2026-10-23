"""
Utah Devils 경기 기록지 v3.0 파서 (분석 플랫폼 lib/gameSheet.ts와 같은 규칙).

타석마다 적은 결과 표기(안타·삼진·희생플라이…)에서 직접 계산하고, 기록지 수식 결과(캐시)가 있으면 대조만 한다.
Excel로 다시 저장하지 않은 파일(수식 결과가 비어 있음)도 정확히 읽기 위해서다.
두 파서가 같은 결과를 내는지는 tests/fixtures/game-v3/expected.json으로 양쪽 테스트가 함께 확인한다.
"""
from __future__ import annotations

import datetime as dt
import re
from dataclasses import dataclass, field

import openpyxl
import pandas as pd

from .stats import ERA_BASE_INNINGS, outs_to_ip

PA_RESULTS: dict[str, dict[str, int]] = {
    "안타": {"ab": 1, "hit": 1},
    "2루타": {"ab": 1, "hit": 1, "double": 1},
    "3루타": {"ab": 1, "hit": 1, "triple": 1},
    "홈런": {"ab": 1, "hit": 1, "hr": 1},
    "삼진": {"ab": 1, "so": 1},
    "루킹삼진": {"ab": 1, "so": 1},
    "낫아웃출루": {"ab": 1, "so": 1},
    "땅볼": {"ab": 1},
    "뜬공": {"ab": 1},
    "직선타": {"ab": 1},
    "병살": {"ab": 1},
    "실책출루": {"ab": 1},
    "야수선택": {"ab": 1},
    "볼넷": {"bb": 1},
    "사구": {"hbp": 1},
    "희생플라이": {"sf": 1},
    "희생번트": {"sh": 1},
    "포수방해": {"ci": 1},
}

BAT_COLS = ["order", "name", "position", "pa", "ab", "runs", "hits", "doubles", "triples", "hr", "rbi", "bb", "hbp",
            "so", "sb", "sf", "sh", "ci"]
PIT_COLS = ["name", "decision", "outs", "pitches", "ha", "runs_allowed", "er", "bb", "hbp", "so", "hr_allowed",
            "w", "l", "sv", "hld"]

FILENAME_RE = re.compile(r"(\d{1,2})_(\d{1,2})\s*VS\s*(.+?)\s*경기\s*기록", re.I)


def _norm(v) -> str:
    return re.sub(r"\s+", "", str(v if v is not None else "")).lower()


def _int(v):
    if v is None or v == "":
        return None
    try:
        return int(round(float(str(v).replace(",", ""))))
    except ValueError:
        return None


def _date(v) -> str | None:
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime("%Y-%m-%d")
    m = re.search(r"(20\d{2})[.\-/]\s*(\d{1,2})[.\-/]\s*(\d{1,2})", str(v or ""))
    return f"{m.group(1)}-{int(m.group(2)):02d}-{int(m.group(3)):02d}" if m else None


@dataclass
class GameSheet:
    source: str
    info: dict
    batting: pd.DataFrame
    pitching: pd.DataFrame
    highlights: list[str]
    review_status: str | None
    review_pending: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    @property
    def is_reviewed(self) -> bool:
        return self.review_status == "검수 완료"


def is_game_workbook(sheetnames: list[str]) -> bool:
    names = {_norm(n) for n in sheetnames}
    return "타자기록" in names and "투수기록" in names


def _sheet(wb, *names):
    wanted = {_norm(n) for n in names}
    return next((wb[n] for n in wb.sheetnames if _norm(n) in wanted), None)


def _parse_info(ws, ws_formula, warnings) -> dict:
    info = {"date": None, "opponent": None, "is_home": None, "location": None, "score_us": None, "score_them": None,
            "defense_outs": None, "era_base": None, "version": None}
    if ws is None:
        return info
    rows = {}
    for r in range(1, ws.max_row + 1):
        label = ws.cell(r, 2).value
        if label is not None:
            rows[_norm(label)] = r
    at = lambda label: ws.cell(rows[_norm(label)], 3).value if _norm(label) in rows else None  # noqa: E731

    info["date"] = _date(at("날짜"))
    if ws_formula is not None and _norm("날짜") in rows:
        f = ws_formula.cell(rows[_norm("날짜")], 3).value
        if isinstance(f, str) and "TODAY()" in f.upper():
            warnings.append("경기 정보 날짜가 입력되지 않았습니다 (=TODAY() 수식 그대로). 파일을 저장한 날짜가 경기 날짜로 들어갑니다.")
    info["opponent"] = (str(at("상대팀")).strip() or None) if at("상대팀") else None
    ha = str(at("홈/원정") or "").strip()
    info["is_home"] = True if ha == "홈" else False if ha == "원정" else None
    info["location"] = (str(at("장소")).strip() or None) if at("장소") else None
    info["score_us"], info["score_them"] = _int(at("우리 점수")), _int(at("상대 점수"))
    d = at("우리 수비 이닝")
    if d not in (None, ""):
        try:
            v = float(d)
            full, digit = int(v), round((v - int(v)) * 10)
            if digit in (0, 1, 2) and abs((v - full) * 10 - digit) < 1e-6:
                info["defense_outs"] = full * 3 + digit
            else:
                raise ValueError
        except ValueError:
            warnings.append(f"경기 정보 '우리 수비 이닝' 값 {d}은 야구 표기가 아닙니다 (7, 6.1, 6.2 형식).")
    base = at("ERA 기준 이닝")
    info["era_base"] = float(base) if isinstance(base, (int, float)) and base > 0 else None
    info["version"] = str(at("버전")).strip() if at("버전") else None
    return info


def _parse_review(ws):
    status, pending = None, []
    if ws is None:
        return status, pending
    for r in range(1, ws.max_row + 1):
        label, state = ws.cell(r, 2).value, ws.cell(r, 6).value
        if not label:
            continue
        state = str(state).strip() if state is not None else ""
        if _norm(label) == _norm("검수 요약"):
            status = state or None
        elif state in ("확인 필요", "입력 대기"):
            pending.append(f"{label}({state})")
    return status, pending


def _parse_batting(ws, warnings) -> pd.DataFrame:
    header = [c.value for c in ws[1]]
    pa_cols = [i for i, h in enumerate(header) if re.fullmatch(r"\d+타석", _norm(h))]
    col = lambda *names: next((i for i, h in enumerate(header) if _norm(h) in {_norm(n) for n in names}), None)  # noqa: E731
    c = {k: col(*v) for k, v in {
        "order": ["타순"], "name": ["이름"], "position": ["포지션"], "runs": ["득점"], "rbi": ["타점"], "sb": ["도루"],
        "c_pa": ["타석 PA"], "c_ab": ["타수"], "c_hits": ["안타"], "c_so": ["삼진"], "c_bb": ["볼넷"],
    }.items()}
    if not pa_cols:
        raise ValueError("타자 기록에 '1타석' 같은 타석 칸이 없습니다 — 경기 기록지 v3.0 형식이 아닙니다.")

    out = []
    for r, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        get = lambda i: row[i] if i is not None and i < len(row) else None  # noqa: E731
        name = str(get(c["name"]) or "").strip()
        codes = [str(row[i]).strip() for i in pa_cols if i < len(row) and row[i] not in (None, "")]
        extras = [_int(get(c[k])) or 0 for k in ("runs", "rbi", "sb")]
        if "합계" in name:
            break
        if not name:
            if codes or any(extras):
                warnings.append(f"타자 기록 {r}행: 이름 없이 기록이 입력됨 — 건너뜀")
            continue
        t = dict.fromkeys(["pa", "ab", "hit", "double", "triple", "hr", "so", "bb", "hbp", "sf", "sh", "ci"], 0)
        for code in codes:
            t["pa"] += 1
            effect = PA_RESULTS.get(re.sub(r"\s+", "", code))
            if effect is None:
                warnings.append(f"타자 기록 {r}행 {name}: 목록에 없는 타석 결과 '{code}' — 타석으로만 셈")
                continue
            for k, v in effect.items():
                t[k] += v
        cached = [("타석", t["pa"], _int(get(c["c_pa"]))), ("타수", t["ab"], _int(get(c["c_ab"]))),
                  ("안타", t["hit"], _int(get(c["c_hits"]))), ("삼진", t["so"], _int(get(c["c_so"]))),
                  ("볼넷", t["bb"], _int(get(c["c_bb"])))]
        diff = [f"{k} {s}→{m}" for k, m, s in cached if s is not None and s != m]
        if diff:
            warnings.append(f"타자 기록 {name}: 기록지 수식 값과 타석 결과 계산이 다름 ({', '.join(diff)}) — 타석 결과 기준으로 저장")
        out.append({
            "order": _int(get(c["order"])) or len(out) + 1, "name": name, "position": str(get(c["position"]) or "").strip(),
            "pa": t["pa"], "ab": t["ab"], "runs": extras[0], "hits": t["hit"], "doubles": t["double"],
            "triples": t["triple"], "hr": t["hr"], "rbi": extras[1], "bb": t["bb"], "hbp": t["hbp"], "so": t["so"],
            "sb": extras[2], "sf": t["sf"], "sh": t["sh"], "ci": t["ci"],
        })
    return pd.DataFrame(out, columns=BAT_COLS)


def _parse_pitching(ws, warnings) -> pd.DataFrame:
    header = [c.value for c in ws[1]]
    col = lambda *names: next((i for i, h in enumerate(header) if _norm(h) in {_norm(n) for n in names}), None)  # noqa: E731
    c = {k: col(v) for k, v in {"name": "이름", "decision": "결과", "outs": "아웃", "pitches": "투구수", "ha": "피안타",
                                 "runs": "실점", "er": "자책", "bb": "볼넷", "hbp": "사구", "so": "삼진", "hr": "피홈런"}.items()}
    if c["outs"] is None:
        raise ValueError("투수 기록에 '아웃' 칸이 없습니다 — 경기 기록지 v3.0 형식이 아닙니다.")
    out = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        get = lambda i: row[i] if i is not None and i < len(row) else None  # noqa: E731
        name = str(get(c["name"]) or "").strip()
        if "합계" in name:
            break
        if not name:
            continue
        outs = _int(get(c["outs"]))
        if outs is None:
            warnings.append(f"투수 기록 {name}: 아웃 수가 비어 있음 — 0이닝으로 저장")
        er = _int(get(c["er"]))
        if er is None:
            warnings.append(f"투수 기록 {name}: 자책점 미확정(빈칸) — 0으로 저장되며 ERA에 반영되지 않음")
        decision = str(get(c["decision"]) or "").strip()
        out.append({
            "name": name, "decision": decision, "outs": outs or 0, "pitches": _int(get(c["pitches"])) or 0,
            "ha": _int(get(c["ha"])) or 0, "runs_allowed": _int(get(c["runs"])) or 0, "er": er or 0,
            "bb": _int(get(c["bb"])) or 0, "hbp": _int(get(c["hbp"])) or 0, "so": _int(get(c["so"])) or 0,
            "hr_allowed": _int(get(c["hr"])) or 0,
            "w": int(decision == "승"), "l": int(decision == "패"), "sv": int(decision == "세"), "hld": int(decision == "홀"),
        })
    return pd.DataFrame(out, columns=PIT_COLS)


def parse_game_workbook(path_or_buffer, filename: str, fallback_season: str | None = None) -> GameSheet:
    if hasattr(path_or_buffer, "seek"):
        path_or_buffer.seek(0)
    wb = openpyxl.load_workbook(path_or_buffer, data_only=True)
    if hasattr(path_or_buffer, "seek"):
        path_or_buffer.seek(0)
    wb_f = openpyxl.load_workbook(path_or_buffer, data_only=False)
    warnings: list[str] = []

    info = _parse_info(_sheet(wb, "경기 정보"), _sheet(wb_f, "경기 정보"), warnings)
    status, pending = _parse_review(_sheet(wb, "검수"))

    m = FILENAME_RE.search(filename)
    if m:
        year = (info["date"] or "")[:4] or fallback_season or str(dt.date.today().year)
        fn_date = f"{year}-{int(m.group(1)):02d}-{int(m.group(2)):02d}"
        fn_opp = m.group(3).strip()
        if info["date"] and fn_date != info["date"]:
            warnings.append(f"파일명 날짜({fn_date})와 경기 정보 날짜({info['date']})가 다릅니다 — 경기 정보 날짜로 저장합니다.")
        if info["opponent"] and _norm(fn_opp) != _norm(info["opponent"]):
            warnings.append(f"파일명 상대팀({fn_opp})과 경기 정보 상대팀({info['opponent']})이 다릅니다 — 경기 정보 기준으로 저장합니다.")
        if not info["date"]:
            warnings.append(f"경기 정보 시트에 날짜가 없어 파일명 날짜({fn_date})로 저장합니다.")
        info["date"] = info["date"] or fn_date
        info["opponent"] = info["opponent"] or fn_opp
    info["season"] = info["date"][:4] if info["date"] else fallback_season

    batting = _parse_batting(_sheet(wb, "타자 기록"), warnings)
    pitching = _parse_pitching(_sheet(wb, "투수 기록"), warnings)
    pitching["ip"] = pitching["outs"].map(outs_to_ip)

    hl_ws = _sheet(wb, "주요 기록")
    highlights = []
    if hl_ws is not None:
        for label, value, *_ in hl_ws.iter_rows(values_only=True):
            if label and value not in (None, "", "-"):
                highlights.append(f"{str(label).strip()}: {str(value).strip()}")

    if status and status != "검수 완료":
        warnings.append(f"검수 시트 상태: {status}" + (f" — {', '.join(pending)}" if pending else ""))
    elif not status:
        warnings.append("검수 시트를 찾지 못했거나 검수 요약이 비어 있습니다.")
    runs = int(batting["runs"].sum())
    if info["score_us"] is not None and runs != info["score_us"]:
        warnings.append(f"우리 점수 {info['score_us']} ≠ 타자 득점 합 {runs}")
    allowed = int(pitching["runs_allowed"].sum())
    if info["score_them"] is not None and allowed != info["score_them"]:
        warnings.append(f"상대 점수 {info['score_them']} ≠ 투수 실점 합 {allowed}")
    outs = int(pitching["outs"].sum())
    if info["defense_outs"] is not None and outs != info["defense_outs"]:
        warnings.append(f"수비 아웃 {info['defense_outs']} ≠ 투수 아웃 합 {outs}")
    if info["era_base"] is not None and info["era_base"] != ERA_BASE_INNINGS:
        warnings.append(f"기록지 ERA 기준 이닝이 {info['era_base']:g}입니다. 콘솔은 {ERA_BASE_INNINGS}이닝 기준으로 계산합니다.")

    return GameSheet(source=filename, info=info, batting=batting, pitching=pitching, highlights=highlights,
                     review_status=status, review_pending=pending, warnings=warnings)
