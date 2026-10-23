"""
Utah Devils 기록지(엑셀) 파서 — 기존 TypeScript 업로드 파이프라인
(app/api/upload/route.ts + lib/spreadsheetImport.ts)의 Python 이식본.

지원 시트 형식
  · 시즌 요약   "2023", "Career Totals"           : 타자 블록 → 투수 블록이 위아래로 쌓인 구조
  · 학기 요약   "23 Spring", "2024 Fall"
  · 경기별 기록 "23 Spring (타자)", "24 Fall (투수)" : "7. 홍길동" 라벨 아래 선수별 경기 로그
  · 무시        "Hall of Fame" 등

기존 구현과 달라진 점
  · 컬럼 매칭은 '정확 일치 → 포함 일치' 순서. (기존: 포함 일치만 → 'WHIP'이 'ip' 별칭에,
    'K/9'가 'k' 별칭에 걸릴 수 있었고 헤더 순서 덕분에 우연히 맞았음)
  · 이닝은 아웃카운트(정수)로 변환해 보관. 표기 오류(13.5 등)는 이슈로 보고.
  · 파싱과 DB 쓰기를 분리. 파서는 순수 함수라 테스트 가능.
"""
from __future__ import annotations

import datetime as dt
import re
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable

import openpyxl
import pandas as pd

from .stats import BATTING_COUNTS, PITCHING_COUNTS, InvalidInningsError, ip_to_outs

# ── 별칭 ─────────────────────────────────────────────────────────────────────
COLUMN_ALIASES: dict[str, list[str]] = {
    "number": ["배번", "번호", "등번호", "선수번호", "no", "num", "number", "jersey"],
    "name": ["이름", "선수명", "성명", "player", "playername", "name"],
    "date": ["날짜", "경기일", "경기날짜", "일자", "date", "gamedate"],
    "opponent": ["상대팀", "상대", "opponent", "vs"],
    "g": ["경기", "g", "games"],
    "pa": ["타석", "pa"],
    "ab": ["타수", "ab"],
    "runs": ["득점", "r", "runs"],
    "hits": ["안타", "h", "hits"],
    "doubles": ["2루타", "2b", "doubles"],
    "triples": ["3루타", "3b", "triples"],
    "hr": ["홈런", "hr"],
    "rbi": ["타점", "rbi"],
    "bb": ["볼넷", "4구", "bb", "walks"],
    "hbp": ["사구", "hbp"],
    "so": ["삼진", "so", "k"],
    "sb": ["도루", "sb"],
    "w": ["승", "w"],
    "l": ["패", "l"],
    "sv": ["세", "세이브", "sv"],
    "hld": ["홀", "홀드", "hld"],
    "ip": ["이닝", "ip"],
    "ha": ["피안타", "ha"],
    "runs_allowed": ["실점", "ra"],
    "er": ["자책", "자책점", "er"],
    "hr_allowed": ["피홈런", "hra"],
    "bf": ["타자", "bf"],
    # 시트에 계산돼 있는 비율 지표 (검증용)
    "sheet_avg": ["타율", "avg"],
    "sheet_obp": ["출루율", "obp"],
    "sheet_slg": ["장타율", "slg"],
    "sheet_ops": ["ops"],
    "sheet_era": ["era"],
    "sheet_whip": ["whip"],
}
# 투수 블록에서만 쓰는 이름 (타자 블록의 '타수'와 충돌 방지)
PITCHING_ONLY = {"ab_against": ["타수"]}

BATTING_FIELDS = ["number", "name", "date", "opponent", *BATTING_COUNTS, "sheet_avg", "sheet_obp", "sheet_slg", "sheet_ops"]
PITCHING_FIELDS = ["number", "name", "date", "opponent", "g", "w", "l", "sv", "hld", "ip", "ha", "runs_allowed", "er",
                   "bb", "hbp", "so", "hr_allowed", "bf", "ab_against", "sheet_era", "sheet_whip"]

TERMS = {"spring": "Spring", "summer": "Summer", "fall": "Fall", "autumn": "Fall"}


def norm(value) -> str:
    text = unicodedata.normalize("NFKC", str(value if value is not None else "")).strip().lower()
    return re.sub(r"[^0-9a-z가-힣]+", "", text)


def map_columns(header: list, kind: str) -> dict[str, int]:
    """헤더 행 → {필드: 열 인덱스}. 정확 일치를 먼저 잡고, 남은 필드만 포함 일치로 찾는다."""
    fields = BATTING_FIELDS if kind == "batting" else PITCHING_FIELDS
    aliases = {f: COLUMN_ALIASES.get(f, []) for f in fields}
    if kind == "pitching":
        aliases.update(PITCHING_ONLY)
    tokens = [norm(c) for c in header]
    mapping: dict[str, int] = {}
    used: set[int] = set()
    for exact in (True, False):
        for f, al in aliases.items():
            if f in mapping:
                continue
            for i, tok in enumerate(tokens):
                if i in used or not tok:
                    continue
                hit = any(tok == norm(a) for a in al) if exact else any(len(norm(a)) > 1 and norm(a) in tok for a in al)
                if hit:
                    mapping[f] = i
                    used.add(i)
                    break
    return mapping


def detect_kind(header: list) -> str | None:
    toks = {norm(c) for c in header}
    if toks & {"이닝", "ip", "자책", "er"}:
        return "pitching"
    if toks & {"타석", "pa", "타율", "avg"}:
        return "batting"
    return None


# ── 결과 모델 ────────────────────────────────────────────────────────────────
@dataclass
class Issue:
    level: str  # "error" | "warning" | "info"
    sheet: str
    message: str
    row: int | None = None  # 엑셀 행 번호(1부터)

    def as_dict(self):
        return {"수준": self.level, "시트": self.sheet, "행": self.row, "내용": self.message}


@dataclass
class ParsedWorkbook:
    source: str
    season_fallback: str | None
    # scope 컬럼: "2023" / "2023 Spring" / "Career"
    batting: pd.DataFrame = field(default_factory=pd.DataFrame)
    pitching: pd.DataFrame = field(default_factory=pd.DataFrame)
    # 경기별 로그
    batting_games: pd.DataFrame = field(default_factory=pd.DataFrame)
    pitching_games: pd.DataFrame = field(default_factory=pd.DataFrame)
    # 시트 하단 합계 행(통산 / "2023 Spring 시즌")
    footers: pd.DataFrame = field(default_factory=pd.DataFrame)
    issues: list[Issue] = field(default_factory=list)
    sheets: list[dict] = field(default_factory=list)

    @property
    def scopes(self) -> list[str]:
        vals = set(self.batting.get("scope", pd.Series(dtype=str))) | set(self.pitching.get("scope", pd.Series(dtype=str)))
        return sorted(vals)

    def issues_frame(self) -> pd.DataFrame:
        return pd.DataFrame([i.as_dict() for i in self.issues], columns=["수준", "시트", "행", "내용"])


# ── 시트 분류 ────────────────────────────────────────────────────────────────
def season_from_filename(name: str) -> str | None:
    m = re.search(r"(20\d{2})", name)
    if m:
        return m.group(1)
    if re.search(r"career", name, re.I):
        return "Career"
    return None


def classify_sheet(sheet_name: str, fallback_season: str | None) -> dict:
    raw = sheet_name.strip()
    low = raw.lower()
    game_kind = None
    if re.search(r"\((타자|타격)\)", raw):
        game_kind = "batting"
    elif re.search(r"\(투수\)", raw):
        game_kind = "pitching"

    season = None
    m = re.search(r"\b(20\d{2})\b", raw)
    if m:
        season = m.group(1)
    else:
        m = re.match(r"^(\d{2})(?:\s|$)", raw)
        if m:
            season = f"20{m.group(1)}"
    term = next((v for k, v in TERMS.items() if k in low), None)

    if "hall of fame" in low or "명예의전당" in norm(raw):
        return {"type": "ignore", "reason": "명예의 전당 시트"}
    if game_kind:
        return {"type": "games", "kind": game_kind, "season": season or fallback_season, "term": term}
    if "career" in low or "통산" in raw:
        return {"type": "summary", "season": "Career", "term": None}
    if season and term:
        return {"type": "summary", "season": season, "term": term}
    if season and re.fullmatch(r"(20)?\d{2}", raw):
        return {"type": "summary", "season": season, "term": None}
    return {"type": "unknown"}


# ── 셀 유틸 ──────────────────────────────────────────────────────────────────
def to_int(v) -> int:
    if v is None or v == "":
        return 0
    if isinstance(v, str):
        v = v.strip()
        if not re.fullmatch(r"-?\d+(\.\d+)?", v):
            return 0
    try:
        return int(round(float(v)))
    except (TypeError, ValueError):
        return 0


def to_float(v):
    if v is None or isinstance(v, str) and not re.fullmatch(r"\s*-?\d+(\.\d+)?\s*", v):
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def parse_date(v, season: str | None) -> str | None:
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, (int, float)) and 20000 < v < 80000:  # 엑셀 일련번호
        return (dt.date(1899, 12, 30) + dt.timedelta(days=int(v))).isoformat()
    s = str(v or "").strip()
    m = re.match(r"^(20\d{2})[.\-/ ]\s*(\d{1,2})[.\-/ ]\s*(\d{1,2})", s)
    if m:
        return f"{m.group(1)}-{int(m.group(2)):02d}-{int(m.group(3)):02d}"
    m = re.match(r"^(\d{1,2})[.\-/](\d{1,2})$", s)
    if m and season and season.isdigit():
        return f"{season}-{int(m.group(1)):02d}-{int(m.group(2)):02d}"
    return None


def term_for_date(iso: str | None) -> str | None:
    if not iso:
        return None
    month = int(iso[5:7])
    return "Spring" if month <= 6 else "Summer" if month <= 8 else "Fall"


IDENTITY_RE = re.compile(r"^\s*(\d{1,3})\s*[.\-]\s*(.+?)\s*$")


def clean_name(name) -> str:
    return re.sub(r"\s+", "", str(name or "")).strip()


# ── 블록 파싱 ────────────────────────────────────────────────────────────────
def _stat_row(row: list, cols: dict[str, int], kind: str, sheet: str, excel_row: int, issues: list[Issue]) -> dict:
    get = lambda f: row[cols[f]] if f in cols and cols[f] < len(row) else None  # noqa: E731
    rec: dict = {}
    if kind == "batting":
        for f in BATTING_COUNTS:
            rec[f] = to_int(get(f))
        for f in ("sheet_avg", "sheet_obp", "sheet_slg", "sheet_ops"):
            rec[f] = to_float(get(f))
    else:
        for f in PITCHING_COUNTS:
            if f == "outs":
                raw_ip = get("ip")
                rec["ip_raw"] = to_float(raw_ip)
                try:
                    rec["outs"] = ip_to_outs(rec["ip_raw"])
                except InvalidInningsError as e:
                    rec["outs"] = ip_to_outs(rec["ip_raw"], strict=False)
                    issues.append(Issue("error", sheet, f"{e} — 소수 이닝으로 간주해 {rec['outs']}아웃 처리", excel_row))
            else:
                rec[f] = to_int(get(f))
        for f in ("sheet_era", "sheet_whip"):
            rec[f] = to_float(get(f))
    return rec


def _is_blank(row) -> bool:
    return all(c is None or str(c).strip() == "" for c in row)


def parse_summary_sheet(rows: list[list], sheet: str, scope: str, issues: list[Issue]):
    """타자/투수 블록이 쌓인 요약 시트 → (batting rows, pitching rows, footers)."""
    out = {"batting": [], "pitching": []}
    footers = []
    i = 0
    while i < len(rows):
        header = rows[i]
        toks = [norm(c) for c in header[:3]]
        if not ({"배번", "no", "number"} & set(toks) and {"이름", "name", "선수명"} & set(toks)):
            i += 1
            continue
        kind = detect_kind(header)
        if not kind:
            i += 1
            continue
        cols = map_columns(header, kind)
        j = i + 1
        seen = set()
        while j < len(rows):
            row = rows[j]
            if _is_blank(row):
                break
            first = row[cols["number"]] if cols["number"] < len(row) else None
            if to_float(first) is None:  # '통산', '2023 Spring 시즌' → 합계 행
                # 합계 행은 저장하지 않으므로 이닝 표기 오류는 차단 사유가 아니다.
                # (기록지 SUM이 3.1+3.2를 6.3으로 더하는 구조라 거의 항상 발생)
                footer_issues: list[Issue] = []
                rec = _stat_row(row, cols, kind, sheet, j + 1, footer_issues)
                for fi in footer_issues:
                    issues.append(Issue("info", sheet, f"합계 행 {fi.message.split(' — ')[0]} — 시트 SUM이 이닝 표기를 소수로 더함", fi.row))
                footers.append({"scope": scope, "kind": kind, "label": str(first), "sheet": sheet, "row": j + 1, **rec})
                j += 1
                continue
            number = to_int(first)
            name = clean_name(row[cols["name"]] if cols["name"] < len(row) else "")
            if not name:
                issues.append(Issue("warning", sheet, f"이름이 비어 있는 행 (배번 {number}) — 건너뜀", j + 1))
                j += 1
                continue
            key = (number, name)
            if key in seen:
                issues.append(Issue("error", sheet, f"같은 블록에 중복 선수 {number}. {name}", j + 1))
            seen.add(key)
            rec = _stat_row(row, cols, kind, sheet, j + 1, issues)
            out[kind].append({"scope": scope, "sheet": sheet, "row": j + 1, "number": number, "name": name, **rec})
            j += 1
        i = j
    return out["batting"], out["pitching"], footers


def parse_game_sheet(rows: list[list], sheet: str, kind: str, season: str | None, term: str | None, issues: list[Issue]):
    """선수별 경기 로그 시트 → 경기 단위 행. 출전하지 않은 경기(전부 0)는 버린다."""
    games = []
    current = None
    cols = None
    placeholder_blocks = 0
    i = 0
    while i < len(rows):
        row = rows[i]
        first = row[0] if row else None
        label = IDENTITY_RE.match(str(first)) if isinstance(first, str) else None
        if label and all(c is None for c in row[1:4]):
            number, name = int(label.group(1)), clean_name(label.group(2))
            current = {"number": number, "name": name, "placeholder": number == 0 or name in ("선수", "player")}
            placeholder_blocks += current["placeholder"]
            cols = None
            i += 1
            continue
        if current and norm(first) in ("날짜", "date"):
            cols = map_columns(row, kind)
            i += 1
            continue
        if current and cols and first is not None:
            date = parse_date(first, season)
            if date:
                rec = _stat_row(row, cols, kind, sheet, i + 1, issues)
                played = any(rec.get(f, 0) for f in (BATTING_COUNTS if kind == "batting" else PITCHING_COUNTS))
                if played:
                    if current["placeholder"]:
                        issues.append(Issue("error", sheet, "이름 미입력('0. 선수') 블록에 기록이 있음", i + 1))
                    else:
                        opp = row[cols["opponent"]] if "opponent" in cols else None
                        games.append({
                            "season": season, "term": term or term_for_date(date), "sheet": sheet, "row": i + 1,
                            "number": current["number"], "name": current["name"], "date": date,
                            "opponent": str(opp).strip() if opp else None, **rec,
                        })
        i += 1
    if placeholder_blocks:
        issues.append(Issue("info", sheet, f"이름 미입력 템플릿 블록 {placeholder_blocks}개 (기록 없음) — 무시"))
    if not term:
        issues.append(Issue("warning", sheet, "시트 이름에 학기가 없어 경기 날짜(월)로 학기를 추정함"))
    return games


# ── 진입점 ───────────────────────────────────────────────────────────────────
def parse_workbook(path_or_buffer, filename: str | None = None) -> ParsedWorkbook:
    filename = filename or (Path(path_or_buffer).name if isinstance(path_or_buffer, (str, Path)) else "upload.xlsx")
    fallback = season_from_filename(filename)
    wb = openpyxl.load_workbook(path_or_buffer, data_only=True, read_only=True)
    result = ParsedWorkbook(source=filename, season_fallback=fallback)
    bat, pit, bat_g, pit_g, foot = [], [], [], [], []

    for ws in wb.worksheets:
        info = classify_sheet(ws.title, fallback)
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        result.sheets.append({"sheet": ws.title, **info})
        if info["type"] == "summary":
            scope = info["season"] if not info.get("term") else f"{info['season']} {info['term']}"
            b, p, f = parse_summary_sheet(rows, ws.title, scope, result.issues)
            bat += b
            pit += p
            foot += f
            if not b and not p:
                result.issues.append(Issue("warning", ws.title, "요약 시트로 보이지만 기록 블록을 찾지 못함"))
        elif info["type"] == "games":
            g = parse_game_sheet(rows, ws.title, info["kind"], info["season"], info["term"], result.issues)
            (bat_g if info["kind"] == "batting" else pit_g).extend(g)
        elif info["type"] == "ignore":
            result.issues.append(Issue("info", ws.title, f"건너뜀: {info['reason']}"))
        else:
            result.issues.append(Issue("warning", ws.title, "형식을 알 수 없는 시트 — 건너뜀"))
    wb.close()

    result.batting = pd.DataFrame(bat)
    result.pitching = pd.DataFrame(pit)
    result.batting_games = pd.DataFrame(bat_g)
    result.pitching_games = pd.DataFrame(pit_g)
    result.footers = pd.DataFrame(foot)
    return result


def parse_many(paths: Iterable[Path]) -> list[ParsedWorkbook]:
    return [parse_workbook(p) for p in paths]
